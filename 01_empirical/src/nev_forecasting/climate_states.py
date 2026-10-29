"""Lagged daily CCPU with training-window quartiles frozen for each test window."""
from __future__ import annotations
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from .common import resolve


def market_calendar(config):
    baseline = pd.read_parquet(resolve(config["paths"]["baseline"]), columns=["datetime"])
    return pd.DatetimeIndex(sorted(pd.to_datetime(baseline.datetime).unique()))

def load_ccpu(config: dict) -> pd.DataFrame:
    path = resolve(config["paths"]["ccpu"])
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet_name = config["ccpu"]["sheet"]
    if sheet_name not in workbook.sheetnames:
        raise AssertionError(f"CCPU workbook does not contain sheet {sheet_name}")
    rows = list(workbook[sheet_name].iter_rows(values_only=True))
    if rows[0][:2] != ("date", "CCPU"):
        raise AssertionError(f"Unexpected CCPU header: {rows[0]}")
    frame = pd.DataFrame(rows[1:], columns=["date", "CCPU"])
    frame["date"] = pd.to_datetime(frame["date"]).dt.normalize()
    frame["CCPU"] = pd.to_numeric(frame["CCPU"], errors="raise")
    if len(frame) != 9497:
        raise AssertionError(f"Expected 9,497 CCPU rows, found {len(frame):,}")
    if frame.isna().any().any() or frame["date"].duplicated().any():
        raise AssertionError("CCPU contains missing or duplicate values")
    if not np.isfinite(frame["CCPU"]).all():
        raise AssertionError("CCPU contains nonfinite values")
    expected_start, expected_end = pd.Timestamp("2000-01-01"), pd.Timestamp("2025-12-31")
    if frame["date"].min() != expected_start or frame["date"].max() != expected_end:
        raise AssertionError("Unexpected CCPU coverage")
    return frame.sort_values("date").reset_index(drop=True)


def window_boundaries(config: dict) -> pd.DataFrame:
    metrics = pd.read_csv(resolve(config["paths"]["window_metrics"])).copy()
    metrics = metrics[(metrics["Track"] == "main") & (metrics["Model"] == "XGBoost") & (metrics["Features"] == "tech")]
    columns = ["Window", "TrainStart", "TrainEnd", "TestStart", "TestEnd"]
    windows = metrics[columns].drop_duplicates().sort_values("Window")
    if len(windows) != 10:
        raise AssertionError(f"Expected 10 rolling windows, found {len(windows)}")
    for column in columns[1:]:
        windows[column] = pd.to_datetime(windows[column]).dt.normalize()
    return windows


def add_ccpu_states(
    panel: pd.DataFrame,
    config: dict,
    lag_days: int,
    low_quantile: float,
    high_quantile: float,
    prefix: str,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    ccpu = load_ccpu(config).set_index("date")["CCPU"]
    calendar = market_calendar(config)
    calendar_frame = pd.DataFrame({"datetime": calendar})
    calendar_frame[f"{prefix}_date"] = calendar_frame["datetime"] - pd.to_timedelta(lag_days, unit="D")
    calendar_frame[f"{prefix}_value"] = calendar_frame[f"{prefix}_date"].map(ccpu)
    if calendar_frame[f"{prefix}_value"].isna().any():
        raise AssertionError(f"Missing aligned CCPU values for {prefix}")

    windows = window_boundaries(config)
    thresholds: list[dict] = []
    for row in windows.itertuples(index=False):
        train = calendar_frame[calendar_frame["datetime"].between(row.TrainStart, row.TrainEnd)]
        values = train[f"{prefix}_value"]
        thresholds.append(
            {
                "window_id": int(row.Window),
                "lag_days": int(lag_days),
                "low_quantile": float(low_quantile),
                "high_quantile": float(high_quantile),
                "low_threshold": float(values.quantile(low_quantile)),
                "high_threshold": float(values.quantile(high_quantile)),
                "train_start": row.TrainStart,
                "train_end": row.TrainEnd,
                "training_dates": len(values),
            }
        )
    threshold_frame = pd.DataFrame(thresholds)
    output = panel.merge(
        calendar_frame[["datetime", f"{prefix}_date", f"{prefix}_value"]],
        on="datetime",
        how="left",
        validate="many_to_one",
    ).merge(threshold_frame, on="window_id", how="left", validate="many_to_one")
    state = np.full(len(output), "middle", dtype=object)
    state[output[f"{prefix}_value"].to_numpy() <= output["low_threshold"].to_numpy()] = "low"
    state[output[f"{prefix}_value"].to_numpy() >= output["high_threshold"].to_numpy()] = "high"
    output[f"{prefix}_state"] = state
    if output[f"{prefix}_value"].isna().any():
        raise AssertionError(f"Prediction panel has missing CCPU values for {prefix}")
    return output, threshold_frame
